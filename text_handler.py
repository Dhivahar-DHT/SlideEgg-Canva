from pptx.enum.text import PP_ALIGN

class TextHandler:
    def __init__(self, color_handler):
        self.color_handler = color_handler
        
    def get_text_properties(self, shape):
        """Extract all text properties from a shape"""
        text_props = {
            'paragraphs': []
        }
        
        if not hasattr(shape, 'text_frame'):
            return text_props
            
        try:
            for paragraph in shape.text_frame.paragraphs:
                para_props = self._get_paragraph_properties(paragraph)
                if para_props:
                    text_props['paragraphs'].append(para_props)
                    
        except Exception as e:
            print(f"Error getting text properties: {e}")
            
        return text_props
    
    def _get_paragraph_properties(self, paragraph):
        """Extract properties from a paragraph"""
        try:
            para_props = {
                'text': paragraph.text,
                'align': self._get_alignment(paragraph),
                'spacing_before': paragraph.space_before.pt if hasattr(paragraph, 'space_before') and paragraph.space_before else 0,
                'spacing_after': paragraph.space_after.pt if hasattr(paragraph, 'space_after') and paragraph.space_after else 0,
                'line_spacing': paragraph.line_spacing if hasattr(paragraph, 'line_spacing') else 1.0,
                'runs': []
            }
            
            # Process each run (segment with consistent formatting)
            for run in paragraph.runs:
                run_props = self._get_run_properties(run)
                if run_props:
                    para_props['runs'].append(run_props)
            
            return para_props
            
        except Exception as e:
            print(f"Error getting paragraph properties: {e}")
            return None
    
    def _get_run_properties(self, run):
        """Extract properties from a run"""
        try:
            props = {
                'text': run.text,
                'font': {
                    'name': run.font.name if hasattr(run.font, 'name') else 'Arial',
                    'size': run.font.size.pt if hasattr(run.font, 'size') and run.font.size else 12,
                    'bold': run.font.bold if hasattr(run.font, 'bold') else False,
                    'italic': run.font.italic if hasattr(run.font, 'italic') else False,
                    'underline': run.font.underline if hasattr(run.font, 'underline') else False,
                    'color': self.color_handler.get_text_color(run),
                    'strike': run.font.strike if hasattr(run.font, 'strike') else False,
                    'subscript': run.font.subscript if hasattr(run.font, 'subscript') else False,
                    'superscript': run.font.superscript if hasattr(run.font, 'superscript') else False
                }
            }
            
            return props
            
        except Exception as e:
            print(f"Error getting run properties: {e}")
            return None
    
    def _get_alignment(self, paragraph):
        """Get paragraph alignment"""
        try:
            if hasattr(paragraph, 'alignment'):
                align_map = {
                    PP_ALIGN.LEFT: 'left',
                    PP_ALIGN.CENTER: 'center',
                    PP_ALIGN.RIGHT: 'right',
                    PP_ALIGN.JUSTIFY: 'justify'
                }
                return align_map.get(paragraph.alignment, 'left')
        except Exception as e:
            print(f"Error getting alignment: {e}")
        return 'left'  # Default alignment
    
    def convert_to_fabric_text(self, text_props):
        """Convert text properties to Fabric.js format"""
        try:
            print(f"Converting text properties: {text_props}")  # Debug log
            
            fabric_text = {
                'type': 'textbox',
                'text': '',
                'styles': {},
                'fontSize': 12,  # Default font size
                'fontFamily': 'Arial',  # Default font family
                'textAlign': 'left',  # Default alignment
                'fill': '#000000',  # Default color
                'originX': 'left',
                'originY': 'top'
            }
            
            # Track the most common font properties to set as default
            font_sizes = []
            font_families = []
            colors = []
            alignments = []
            
            current_index = 0
            for paragraph in text_props['paragraphs']:
                print(f"Processing paragraph: {paragraph}")  # Debug log
                
                alignments.append(paragraph['align'])
                
                for run in paragraph['runs']:
                    text = run['text']
                    if text:
                        fabric_text['text'] += text
                        
                        font_sizes.append(run['font']['size'])
                        font_families.append(run['font']['name'])
                        colors.append(run['font']['color'])
                        
                        # Add style information for this run
                        style = {
                            'fontFamily': run['font']['name'],
                            'fontSize': run['font']['size'],
                            'fontWeight': 'bold' if run['font']['bold'] else 'normal',
                            'fontStyle': 'italic' if run['font']['italic'] else 'normal',
                            'underline': run['font']['underline'],
                            'fill': run['font']['color'],
                            'textAlign': paragraph['align']
                        }
                        
                        print(f"Adding style for text '{text}': {style}")  # Debug log
                        
                        # Add style for each character in the run
                        for i in range(len(text)):
                            fabric_text['styles'][str(current_index + i)] = style
                        
                        current_index += len(text)
                
                # Add newline between paragraphs
                if paragraph != text_props['paragraphs'][-1]:
                    fabric_text['text'] += '\n'
                    current_index += 1
            
            # Set the most common properties as defaults
            if font_sizes:
                fabric_text['fontSize'] = max(set(font_sizes), key=font_sizes.count)
            if font_families:
                fabric_text['fontFamily'] = max(set(font_families), key=font_families.count)
            if colors:
                fabric_text['fill'] = max(set(colors), key=colors.count)
            if alignments:
                fabric_text['textAlign'] = max(set(alignments), key=alignments.count)
            
            print(f"Final fabric text object: {fabric_text}")  # Debug log
            return fabric_text
            
        except Exception as e:
            print(f"Error converting to Fabric text: {e}")
            print(f"Text properties: {text_props}")  # Debug log
            return None 